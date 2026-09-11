import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    """Sets background color for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Sets subtle light gray borders for table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def add_callout(doc, text, title="", border_color="1A365D", bg_color="F0F4F8"):
    """Creates a stylized callout box with a thick left border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    if title:
        run_title = p.add_run(f"📌 {title}\n")
        run_title.bold = True
        run_title.font.name = 'Arial'
        run_title.font.size = Pt(11)
        run_title.font.color.rgb = RGBColor(26, 54, 93)
    
    run_text = p.add_run(text)
    run_text.font.name = 'Arial'
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(45, 55, 72)

def generate_doc():
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(45, 55, 72)
    
    # --- HEADER / PORTADA ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(2)
    t_run = title_p.add_run("PROPUESTA DE DESARROLLO E IMPLEMENTACIÓN TECNOLÓGICA")
    t_run.bold = True
    t_run.font.name = 'Arial'
    t_run.font.size = Pt(18)
    t_run.font.color.rgb = RGBColor(26, 54, 93) # Deep Navy
    
    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(16)
    sub_run = subtitle_p.add_run("Sistema Centralizado de Inventario y Control de Flota Vehicular")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(12.5)
    sub_run.font.color.rgb = RGBColor(43, 108, 176)
    
    # Header metadata table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        [("Cliente:", " Sindicato de Choferes Profesionales"), ("Fecha de Emisión:", " Septiembre 2026")],
        [("Proyecto:", " Sistema de Gestión de Activos y Flota"), ("Validez de Oferta:", " 30 Días Calendario")]
    ]
    
    col_widths = [Inches(3.25), Inches(3.25)]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F7FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lbl, val = meta_data[r_idx][c_idx]
            r1 = p.add_run(lbl)
            r1.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = RGBColor(26, 54, 93)
            r2 = p.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = RGBColor(45, 55, 72)
            
    set_table_borders(meta_table, color="CBD5E0", sz="4")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # --- 1. RESUMEN EJECUTIVO ---
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("1. Resumen del Proyecto y Alcance Técnico")
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    p_desc = doc.add_paragraph(
        "El presente documento detalla la propuesta técnico-económica para la puesta en marcha del Sistema Centralizado "
        "de Inventario y Control de Flota Vehicular. La plataforma opera bajo un modelo SaaS (Software como Servicio) "
        "donde el sistema es alojado en servidores en la nube y entregado al Sindicato mediante un enlace (Link) de acceso web "
        "directo, sin necesidad de realizar instalaciones físicas en los equipos de la institución."
    )
    p_desc.paragraph_format.space_after = Pt(8)
    
    add_callout(
        doc,
        "Modalidad de Entrega: Sistema 100% Nube (SaaS). Se proporciona un enlace de acceso web único y seguro. "
        "Compatible con PCs, laptops, tablets y teléfonos celulares sin instalar nada en los dispositivos.\n"
        "Stack: React 19 + TypeScript (Frontend), Python Django REST Framework (API Backend), PostgreSQL 15 (Base de Datos), WebSockets en Tiempo Real.",
        title="Acceso Inmediato sin Instalaciones Locales"
    )
    
    # --- 2. CARACTERISTICAS Y MODULOS ---
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run("2. Descripción Detallada de Módulos, Valor Tecnológico y Alcance")
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    modules = [
        ("Módulo 1: Control de Flota Vehicular & Garita (Módulo CORE - Principal Inversión Tecnológica)",
         "¿Qué hace?: Registro operativo de entradas/salidas en garita, kilometraje inicial/final, combustible, conductor asignado, novedades mecánicas y panel administrativo en vivo.\n"
         "• Justificación de Valor: Es el componente central y de mayor valor técnico del sistema. Implementa arquitectura de WebSockets en tiempo real, validación de inconsistencias en kilometraje e interfaz de garita optimizada para celulares sin recargar la página."),
        
        ("Módulo 2: Captura y Compresión Inteligente de Evidencia Fotográfica",
         "¿Qué hace?: Fotografías in situ tomadas desde el celular en garita para documentar novedades o estado del vehículo.\n"
         "• Justificación de Valor: Motor algorítmico que comprime las imágenes un 90% antes del envío, garantizando velocidad de carga en redes 3G/4G y un ahorro masivo en almacenamiento cloud."),
        
        ("Módulo 3: Gestión Integral de Inventario (Bienes, Muebles y Equipos)",
         "¿Qué hace?: Control de activos fijos institucionales, asignación oficial a custodios por departamento, ubicaciones físicas, estado de conservación y actas de baja.\n"
         "• Justificación de Valor: Garantiza el control patrimonial del Sindicato, eliminando pérdidas de bienes y facilitando auditorías físicas instantáneas."),
        
        ("Módulo 4: Reportes PDF & Actas Automatizadas",
         "¿Qué hace?: Generación automática de actas de entrega-recepción de vehículos y actas de baja en formato PDF oficial.\n"
         "• Justificación de Valor: Emisión de documentos y actas oficiales en PDF que sustituye el uso de carpetas en papel y respalda legalmente las entregas e inventarios."),
        
        ("Módulo 5: Capacitación y Documentación del Sistema",
         "¿Qué hace?: Sesiones de entrenamiento práctico presencial/virtual a guardias y administradores + Manual de Usuario digital.\n"
         "• Justificación de Valor: Garantiza una adopción fluida del software por parte del personal sin curva de aprendizaje prolongada."),
        
        ("Módulo 6: Bitácora de Auditoría e Historial de Alta Seguridad",
         "¿Qué hace?: Registro inmutable de cada acción efectuada en el sistema (usuario, fecha, hora, dirección IP y datos modificados).\n"
         "• Justificación de Valor: Seguridad de nivel bancario que otorga transparencia total ante auditorías internas y directivas.")
    ]
    
    for title, desc in modules:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r_t = p.add_run(f"• {title}\n")
        r_t.bold = True
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = RGBColor(43, 108, 176)
        r_d = p.add_run(desc)
        r_d.font.size = Pt(10)
        r_d.font.color.rgb = RGBColor(45, 55, 72)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    add_callout(
        doc,
        "Si en el futuro el Sindicato requiere la incorporación de nuevos módulos no contemplados en el alcance inicial "
        "(ej. Módulo de Mantenimientos Preventivos Avanzados, Control de Combustible/Gasolineras, Módulo de Sanciones, etc.), "
        "cada nuevo módulo personalizado se cotizará e implementará a partir de $200.00 USD en adelante, "
        "según su complejidad técnica, número de pantallas y reglas de negocio solicitadas.\n"
        "Durante la vigencia del contrato, todo desarrollo adicional se realizará con el desarrollador original.",
        title="Desarrollos Futuros y Tarifas de Módulos Adicionales",
        border_color="D97706",
        bg_color="FEF3C7"
    )
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    
    # --- 3. OPCIONES DE INVERSIÓN Y CONTRATACIÓN ---
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r = h3.add_run("3. Opciones de Inversión y Formas de Pago")
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    p_intro = doc.add_paragraph(
        "A continuación se detallan los esquemas de contratación disponibles para el Sindicato de Choferes Profesionales:"
    )
    p_intro.paragraph_format.space_after = Pt(8)

    # --- OPCION A: LICENCIA DE USO EMPRESARIAL ---
    h_op_a = doc.add_paragraph()
    h_op_a.paragraph_format.space_before = Pt(8)
    h_op_a.paragraph_format.space_after = Pt(4)
    r = h_op_a.add_run("OPCIÓN A: Licencia de Uso Empresarial (Opción Recomendada)")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    p_op_a_desc = doc.add_paragraph(
        "Es la alternativa institucional recomendada. El Sindicato obtiene la licencia de uso y el acceso total al sistema "
        "mediante su enlace web privado, delegando el alojamiento en la nube, el mantenimiento preventivo, "
        "los respaldos continuos de datos y el soporte técnico exclusivo en las manos del desarrollador."
    )
    p_op_a_desc.paragraph_format.space_after = Pt(6)
    
    # Details Box for Option A
    add_callout(
        doc,
        "1. INVERSIÓN INICIAL (PAGO ÚNICO): $2,300.00 USD\n"
        "   • Puesta en Producción en la Nube: Configuración de la infraestructura cloud y generación del Link/Enlace Web de acceso exclusivo.\n"
        "   • Carga Inicial de Datos: Parametrización de vehículos, inventarios de bienes y usuarios autorizados.\n"
        "   • Capacitación del Personal: Sesiones presenciales/virtuales para administradores y guardias de garita.\n"
        "   • Documentación Oficial: Entrega del Manual de Usuario e Instructivo Técnico del Sistema.\n\n"
        "2. PÓLIZA MENSUAL DE ALOJAMIENTO, MANTENIMIENTO Y RESPALDOS: $135.00 USD / mes\n"
        "   • Alojamiento Cloud (Hosting): Garantía de servidor en vivo 24/7 con certificado de seguridad SSL encriptado.\n"
        "   • Mantenimiento Semanal: Monitoreo continuo del servidor, revisión de recursos y optimización de memoria.\n"
        "   • Mantenimiento Quincenal / Mensual: Mantenimiento preventivo de la base de datos PostgreSQL, depuración de logs, optimización de índices y parches de seguridad.\n"
        "   • Respaldos Semanales de Datos: Copias de seguridad automáticas semanales de la base de datos y evidencias fotográficas.\n"
        "   • Respaldos Mensuales Históricos: Archivo acumulativo mensual de seguridad para restauración inmediata ante cualquier contingencia.\n\n"
        "3. PLAZO MÍNIMO DE CONTRATACIÓN: UN (1) AÑO (12 meses de contrato obligatorio).\n\n"
        "4. CLÁUSULA DE EXCLUSIVIDAD TÉCNICA Y MODIFICACIONES:\n"
        "   • El desarrollo, mantenimiento y soporte es de EXCLUSIVIDAD ABSOLUTA del desarrollador original.\n"
        "   • Ningún tercero podrá alterar ni modificar el sistema.\n"
        "   • Cualquier solicitud de nuevos módulos o cambios en el sistema será llamada y contratada EXCLUSIVAMENTE con el desarrollador original.",
        title="Desglose de Servicios y Términos de la Opción A",
        border_color="2563EB",
        bg_color="F8FAFC"
    )
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    # --- OPCION B: TRANSFERENCIA DE PROPIEDAD INTELECTUAL ---
    h_op_b = doc.add_paragraph()
    h_op_b.paragraph_format.space_before = Pt(12)
    h_op_b.paragraph_format.space_after = Pt(4)
    r = h_op_b.add_run("OPCIÓN B: Transferencia Total de Propiedad Intelectual (Código Fuente)")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    p_op_b_desc = doc.add_paragraph(
        "El Sindicato adquiere el 100% de los derechos del software, incluyendo todos los archivos de programación originales "
        "(código fuente en frontend y backend). En esta modalidad, la institución asume la responsabilidad de contratar su propio "
        "servidor y gestionar sus propios desarrolladores a futuro."
    )
    p_op_b_desc.paragraph_format.space_after = Pt(6)

    # Details Box for Option B
    add_callout(
        doc,
        "• COSTO TOTAL DE TRANSFERENCIA: $4,500.00 USD (Pago Único).\n"
        "• INCLUYE: Entrega de repositorios con el código fuente completo, manual de código y capacitación inicial.\n"
        "• DERECHOS: El Sindicato tendrá libertad total para modificar el código o contratar a terceros.\n"
        "• NO INCLUYE: Servidor/hosting, ni respaldos semanales, ni mantenimientos preventivos a futuro, ni cláusula de exclusividad. El cliente asume la infraestructura.",
        title="Condiciones y Términos de la Opción B",
        border_color="64748B",
        bg_color="F1F5F9"
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- 4. TABLA COMPARATIVA DE OPCIONES ---
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r = h4.add_run("4. Cuadro Comparativo Resumen")
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = RGBColor(26, 54, 93)
    
    comp_table = doc.add_table(rows=8, cols=3)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    comp_table.autofit = False
    
    comp_widths = [Inches(2.5), Inches(2.0), Inches(2.0)]
    comp_headers = ["Criterio de Evaluación", "Opción A: Licencia de Uso (Recomendada)", "Opción B: Código Fuente"]
    
    hdr_cells = comp_table.rows[0].cells
    for i, title in enumerate(comp_headers):
        hdr_cells[i].width = comp_widths[i]
        set_cell_background(hdr_cells[i], "1A365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=80, right=80)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)

    comp_rows_data = [
        ("Modalidad de Entrega", "Acceso directo vía Enlace Web (SaaS)", "Entrega de Archivos de Código"),
        ("Inversión Inicial", "$2,300.00 USD (Pago Único)", "$4,500.00 USD (Pago Único)"),
        ("Alojamiento Cloud & Servidor", "Incluido 24/7 en la cuota mensual", "A cargo del Sindicato"),
        ("Frecuencia de Mantenimiento", "Semanal, Quincenal y Mensual", "No incluye"),
        ("Frecuencia de Respaldos", "Semanales y Mensuales Automáticos", "No incluye"),
        ("Capacitación y Documentación", "Incluido en la inversión inicial", "Incluido en la inversión inicial"),
        ("Exclusividad de Desarrollador", "Sí (Contrato 1 año, cambios exclusivos con creador)", "No (Libertad de modificar con terceros)")
    ]

    for r_idx, row_data in enumerate(comp_rows_data):
        row_cells = comp_table.rows[r_idx + 1].cells
        bg_color = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].width = comp_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=80, right=80)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(cell_value)
            r.font.size = Pt(8.5)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(26, 54, 93)
            else:
                r.font.color.rgb = RGBColor(45, 55, 72)

    set_table_borders(comp_table, color="CBD5E0", sz="4")

    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    
    # --- 5. FIRMAS Y ACEPTACIÓN ---
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)
    r = h5.add_run("5. Aceptación y Firmas")
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = RGBColor(26, 54, 93)

    p_sig_text = doc.add_paragraph(
        "En señal de conformidad con los términos, alcance técnico y condiciones comerciales descritas "
        "en el presente informe, las partes firman el presente documento:"
    )
    p_sig_text.paragraph_format.space_after = Pt(24)

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    sig_widths = [Inches(3.25), Inches(3.25)]
    c0 = sig_table.cell(0, 0)
    c1 = sig_table.cell(0, 1)
    
    c0.width = sig_widths[0]
    c1.width = sig_widths[1]
    
    set_cell_margins(c0, top=100, bottom=100, left=50, right=50)
    set_cell_margins(c1, top=100, bottom=100, left=50, right=50)
    
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.add_run("_____________________________________\n").font.color.rgb = RGBColor(160, 174, 192)
    r_s1 = p0.add_run("Desarrollador Tecnológico Exclusivo\n")
    r_s1.bold = True
    r_s1.font.size = Pt(9.5)
    r_s1.font.color.rgb = RGBColor(26, 54, 93)
    p0.add_run("Proveedor de Servicio & Soporte").font.size = Pt(8.5)
    
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.add_run("_____________________________________\n").font.color.rgb = RGBColor(160, 174, 192)
    r_s2 = p1.add_run("Sindicato de Choferes Profesionales\n")
    r_s2.bold = True
    r_s2.font.size = Pt(9.5)
    r_s2.font.color.rgb = RGBColor(26, 54, 93)
    p1.add_run("Representante Legal / Cliente").font.size = Pt(8.5)

    filename = "Informe_Cotizacion_Sistema_Inventario.docx"
    try:
        doc.save(filename)
        print(f"Updated Word Document successfully created: {filename}")
    except PermissionError:
        filename_v2 = "Informe_Cotizacion_Sistema_Inventario_v2.docx"
        doc.save(filename_v2)
        print(f"Primary file was locked by Word. Saved updated file as: {filename_v2}")

if __name__ == "__main__":
    generate_doc()
